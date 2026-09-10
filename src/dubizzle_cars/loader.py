"""Read the named sheet without saving, evaluating formulas, or changing source bytes."""

from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .extraction import extract_metadata
from .models import Car

SHEET_NAME = "cleaned dataset"
COLUMNS = ("Listing_ID", "year", "make", "model", "trim", "title", "description", "photo_url")


class InventoryDataError(ValueError):
    """An input cannot be loaded safely as an inventory."""


def _text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _integer(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = Decimal(str(value))
        if number.is_finite() and number == number.to_integral_value():
            return int(number)
    except InvalidOperation:
        pass
    return None


def load_inventory(path: str | Path) -> tuple[Car, ...]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise InventoryDataError(f"Missing sheet: {SHEET_NAME!r}")
        rows = workbook[SHEET_NAME].iter_rows()
        header = next(rows, ())
        names = [_text(c.value) for c in header]
        if any(names.count(name) != 1 for name in COLUMNS):
            raise InventoryDataError(f"Expected each column exactly once: {COLUMNS}")
        indices = {name: names.index(name) for name in COLUMNS}
        cars, seen = [], set()
        for row_number, row in enumerate(rows, 2):
            cells = {name: row[i] if i < len(row) else None for name, i in indices.items()}
            values = {
                name: cell.value if cell is not None else None for name, cell in cells.items()
            }
            if all(_text(v) is None for v in values.values()):
                continue
            # Reject formulas/errors rather than trust stale Excel cached values.
            if any(c is not None and c.data_type in {"f", "e"} for c in cells.values()):
                raise InventoryDataError(f"Formula or Excel error in row {row_number}")
            raw_id = values["Listing_ID"]
            listing_id = (
                str(_integer(raw_id))
                if isinstance(raw_id, (int, float))
                and not isinstance(raw_id, bool)
                and _integer(raw_id) is not None
                else _text(raw_id)
            )
            if not listing_id or isinstance(raw_id, bool):
                raise InventoryDataError(f"Missing or invalid Listing_ID in row {row_number}")
            if listing_id in seen:
                raise InventoryDataError(f"Duplicate Listing_ID {listing_id!r} in row {row_number}")
            seen.add(listing_id)
            year = _integer(values["year"])
            fields = {name: _text(values[name]) for name in COLUMNS[2:]}
            cars.append(
                Car(
                    listing_id=listing_id,
                    year=year if year is not None and 1000 <= year <= 9999 else None,
                    **fields,
                    metadata=extract_metadata(fields["title"], fields["description"]),
                    source_row=row_number,
                )
            )
        return tuple(cars)
    finally:
        workbook.close()
